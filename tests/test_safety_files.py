"""JV-01: the dangerous-file shield on every ingest path.

Executables, installers and archives are never downloaded, opened or offered
for approval; the owner is warned once. The classifiers are pure; the ingest
paths (userbot, bot, worker) run against real PostgreSQL with Telegram
replaced by stand-ins.
"""

from __future__ import annotations

import io
import zipfile
from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace

import pytest
import sqlalchemy as sa
from telethon.tl.types import User

from miya.bot import handlers, replies
from miya.config import settings
from miya.db import models as m
from miya.db.enums import ChatType, Direction, InteractionSource
from miya.db.session import session_scope
from miya.services import approvals, documents, questions, safety
from miya.services.media_policy import MediaKind, forced_plan, plan_for
from miya.services.persistence import Applied
from miya.userbot import main as userbot
from miya.worker import main as worker
from tests.test_close_and_correct import _Message, bound  # noqa: F401

TZ = settings.tz
CHAT = 515151
PEER = User(id=CHAT, first_name="Akmal", username="akmal_test")
DAY = datetime.now(TZ).replace(hour=14, minute=0, second=0, microsecond=0)


def plan(kind=MediaKind.document, *, filename=None, size=1024, mime=None, **kw):
    return plan_for(
        kind,
        vision_enabled=kw.pop("vision", True),
        docs_enabled=kw.pop("docs", True),
        filename=filename,
        size=size,
        mime=mime,
        **kw,
    )


# --- files on disk -------------------------------------------------------------


def _pdf(text: str) -> bytes:
    """A minimal one-page PDF that pdfplumber reads."""
    content = f"BT /F1 12 Tf 10 50 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 100] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return out


def _docx(path: Path, text: str) -> Path:
    import docx

    document = docx.Document()
    document.add_paragraph(text)
    document.save(path)
    return path


def _xlsx(path: Path, rows: list[list[object]]) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    return path


def _with_entries(path: Path, extra: dict[str, bytes]) -> Path:
    """Copy an Office file, adding entries (a VBA project, an Android manifest)."""
    buffer = io.BytesIO()
    with (
        zipfile.ZipFile(path) as source,
        zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            target.writestr(info, source.read(info.filename))
        for name, data in extra.items():
            target.writestr(name, data)
    path.write_bytes(buffer.getvalue())
    return path


def _windows_program() -> bytes:
    header = bytearray(128)
    header[0:2] = b"MZ"
    header[60:64] = (64).to_bytes(4, "little")
    header[64:68] = b"PE\x00\x00"
    return bytes(header)


# --- normalising names ---------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("foto\uff0eapk", "foto.apk"),
        ("foto\u2024apk", "foto.apk"),
        ("foto\ufe52apk", "foto.apk"),
        ("foto\u3002apk", "foto.apk"),
        ("foto\u2e33apk", "foto.apk"),
        ("Rasm\u202egpj.apk", "rasmgpj.apk"),
        ("ilo\u200bva.a\u200dpk", "ilova.apk"),
        ("Hisob.pdf      .EXE", "hisob.pdf.exe"),
        ("setup.exe. . ", "setup.exe"),
        ("C:\\Users\\x\\yuk.zip", "yuk.zip"),
        ("WhatsApp_Pro.apk/", "whatsapp_pro.apk"),
        ("run.js\\ . ", "run.js"),
        (None, ""),
    ],
)
def test_normalise_filename(raw, expected):
    assert safety.normalise_filename(raw) == expected


# --- classifying before a download ---------------------------------------------


def test_small_apk_is_dangerous_not_unsupported():
    """A small .apk used to fall through to 'unsupported_type' — quietly
    stored, never warned about."""
    result = plan(filename="ilova.apk", size=2_000_000)
    assert result.download is False and result.ask is False
    assert result.skip_reason == "dangerous_file"
    assert result.risk == safety.Verdict("exec", "suffix")


def test_large_apk_never_creates_an_ask():
    """The size branch used to run first, offering a big installer for
    download with two buttons."""
    for size in (settings.doc_max_bytes + 1, settings.media_ask_max_bytes - 1, None):
        result = plan(filename="WhatsApp_Pro.apk", size=size)
        assert result.ask is False, size
        assert result.ask_reason is None
        assert result.skip_reason == "dangerous_file"


def test_the_shield_runs_even_where_documents_are_off():
    assert plan(filename="ilova.apk", docs=False).skip_reason == "dangerous_file"


@pytest.mark.parametrize(
    "name",
    [
        "Rasmlar.jpg.apk",
        "Rasm\u202egpj.apk",
        "foto\uff0eapk",
        "hisob.pdf .exe",
        "hisob.exe.pdf",
        "setup.EXE.",
        "invoice\u200b.apk",
        "rasm.jpg\u2024scr",
        "yuk.zip.pdf",
        "katalog.part1.rar",
        "katalog.7z.001",
        "ilova.xapk",
        "profil.mobileconfig",
        "setup.ex_",
        "setup.ex",
        "WhatsApp_Pro.apk/",
        "run.js\\",
    ],
)
def test_double_extension_and_rlo_names_blocked(name):
    verdict = safety.classify_attachment(name)
    assert verdict.dangerous, name
    assert plan(filename=name).skip_reason == "dangerous_file"
    large = plan(filename=name, size=60_000_000, mime="application/octet-stream")
    assert (large.ask, large.skip_reason) == (False, "dangerous_file"), name


@pytest.mark.parametrize(
    "name",
    [
        "hisob.pdf",
        "gsr.com.pdf",
        "first.class.ticket.pdf",
        "hisob.001.pdf",
        "Akmal aka. Hisob 12.05.2024.pdf",
        "narxlar.xlsx",
        "rasm.jpg",
    ],
)
def test_ordinary_names_pass(name):
    assert safety.classify_attachment(name) == safety.SAFE, name


def test_the_mime_type_counts_without_a_name():
    android = "application/vnd.android.package-archive"
    assert safety.classify_attachment(None, android).kind == "exec"
    assert safety.classify_attachment("rasm.jpg", android).kind == "exec"
    assert safety.classify_attachment("x", "application/x-rar-compressed").kind == (
        "archive"
    )


def test_an_office_file_labelled_zip_is_left_to_its_bytes():
    assert safety.classify_attachment("hisob.docx", "application/zip") == safety.SAFE
    assert safety.classify_attachment("hisob.pdf", "application/zip").kind == "archive"


def test_macro_and_html_kinds():
    assert safety.classify_attachment("hisob.xlsm").kind == "macro"
    assert safety.classify_attachment("hisob.docm").kind == "macro"
    assert safety.classify_attachment("kirish.html").kind == "html"
    assert plan(filename="kirish.html", size=10**9).ask is False
    macro = plan(filename="hisob.xlsm")
    assert macro.read_document is True and macro.risk.kind == "macro"


def test_audio_and_video_names_are_checked_too():
    assert plan(MediaKind.audio, filename="qoshiq.mp3.exe").skip_reason == (
        "dangerous_file"
    )
    assert plan(MediaKind.video, filename="film.apk").ask is False
    assert plan(MediaKind.voice, mime="audio/ogg").download is True


def test_forced_plan_refuses_dangerous():
    for kind, name in (
        (MediaKind.document, "ilova.apk"),
        (MediaKind.document, "arxiv.rar"),
        (MediaKind.video, "film.mp4.exe"),
        (MediaKind.audio, "yozuv.mp3.scr"),
    ):
        result = forced_plan(kind, name)
        assert result.download is False, name
        assert result.skip_reason == "dangerous_file", name
    by_mime = forced_plan(
        MediaKind.document, "fayl", mime="application/vnd.android.package-archive"
    )
    assert by_mime.download is False


def test_the_same_policy_serves_every_channel():
    for channel in ("telegram", "whatsapp", "gmail"):
        result = plan(filename="ilova.apk", channel=channel)
        assert (result.skip_reason, result.channel) == ("dangerous_file", channel)
        assert plan(filename="hisob.pdf", channel=channel).read_document is True
    with pytest.raises(ValueError):
        plan(filename="hisob.pdf", channel="fax")


# --- the bytes of an allowed type ----------------------------------------------


@pytest.mark.parametrize(
    "entry", ["AndroidManifest.xml", "classes.dex", "META-INF/MANIFEST.MF"]
)
def test_ooxml_with_android_entries_rejected(tmp_path, entry):
    path = _with_entries(_docx(tmp_path / "shartnoma.docx", "matn"), {entry: b"x"})
    verdict = safety.sniff_magic(path, ".docx")
    assert verdict == safety.Verdict("exec", "app_package")
    assert verdict.dangerous and not verdict.readable


@pytest.mark.parametrize(
    ("content", "suffix", "kind"),
    [
        (_windows_program(), ".pdf", "exec"),
        (b"\x7fELF\x02\x01\x01" + bytes(64), ".csv", "exec"),
        (b"dex\n035\x00" + bytes(64), ".pdf", "exec"),
        (b"\xcf\xfa\xed\xfe" + bytes(64), ".docx", "exec"),
        (b"Rar!\x1a\x07\x00" + bytes(16), ".pdf", "archive"),
        (b"SZDD\x88\xf0\x27\x33A\x00" + bytes(16), ".pdf", "archive"),
        (b"KWAJ\x88\xf0\x27\xd1" + bytes(16), ".txt", "archive"),
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(64), ".xlsx", "mismatch"),
        (b"not a pdf at all", ".pdf", "mismatch"),
    ],
)
def test_signatures_must_match_the_suffix(tmp_path, content, suffix, kind):
    path = tmp_path / f"fayl{suffix}"
    path.write_bytes(content)
    assert safety.sniff_magic(path, suffix).kind == kind


def test_a_zip_posing_as_a_pdf_is_an_archive(tmp_path):
    path = tmp_path / "hisob.pdf"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("ichida.exe", b"MZ")
    assert safety.sniff_magic(path, ".pdf") == safety.Verdict("archive", "zip_container")


def test_a_zip_bomb_is_refused_by_its_central_directory(tmp_path):
    path = tmp_path / "hisob.pdf"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("pad.bin", bytes(20 * 1024 * 1024))
    assert safety.sniff_magic(path, ".pdf") == safety.Verdict("archive", "zip_bomb")


def test_an_office_file_over_the_zip_limits_is_kept_not_called_an_installer(tmp_path):
    """A spreadsheet past the ratio limit is a big spreadsheet: never parsed,
    but never deleted or reported as an installer either."""
    path = _xlsx(tmp_path / "hisob.xlsx", [["a"]])
    _with_entries(path, {"xl/media/pad.bin": bytes(20 * 1024 * 1024)})
    verdict = safety.sniff_magic(path, ".xlsx")
    assert verdict == safety.Verdict("mismatch", "zip_too_large")
    assert not verdict.dangerous and not verdict.readable


def _big_xlsx(path: Path) -> Path:
    """A real workbook grown past 100 MB uncompressed at an ordinary ratio:
    a second sheet of unique rows, so it does not compress like a bomb."""
    _xlsx(path, [["mijoz", "kod"]])
    with (
        zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED, compresslevel=1) as book,
        book.open("xl/worksheets/sheet2.xml", "w") as sheet,
    ):
        for start in range(0, 1_400_000, 20_000):
            sheet.write(
                "".join(
                    f'<row r="{row}"><c t="inlineStr"><is><t>Mijoz {row:07d} '
                    f"GS{row * 7919 % 10**9:09d}</t></is></c></row>"
                    for row in range(start, start + 20_000)
                ).encode()
            )
    return path


@pytest.fixture(scope="module")
def big_xlsx(tmp_path_factory):
    return _big_xlsx(tmp_path_factory.mktemp("big") / "hisobot.xlsx")


def test_a_big_real_spreadsheet_is_not_dangerous(big_xlsx):
    with zipfile.ZipFile(big_xlsx) as archive:
        infos = archive.infolist()
    total = sum(i.file_size for i in infos)
    packed = sum(i.compress_size for i in infos)
    assert total > safety.ZIP_MAX_UNCOMPRESSED
    assert total < safety.ZIP_MAX_RATIO * packed
    assert big_xlsx.stat().st_size < settings.doc_max_bytes
    verdict = safety.sniff_magic(big_xlsx, ".xlsx")
    assert verdict == safety.Verdict("mismatch", "zip_too_large")
    assert not verdict.dangerous


def test_a_vba_project_makes_a_macro_document(tmp_path):
    path = _with_entries(
        _docx(tmp_path / "x.docx", "matn"), {"word/vbaProject.bin": b"v"}
    )
    assert safety.sniff_magic(path, ".docx") == safety.Verdict("macro", "vba_project")


def test_a_csv_that_starts_with_mz_is_still_a_csv(tmp_path):
    path = tmp_path / "mijozlar.csv"
    path.write_text("MZ Logistics,GS367\nAkmal,GS412\n", encoding="utf-8")
    assert safety.sniff_magic(path, ".csv") == safety.SAFE


def test_pdf_docx_xlsx_still_read(tmp_path):
    pdf = tmp_path / "hisob.pdf"
    pdf.write_bytes(_pdf("Hisob 5000 USD"))
    docx_path = _docx(tmp_path / "shartnoma.docx", "Yuk tashish shartnomasi")
    xlsx_path = _xlsx(tmp_path / "narxlar.xlsx", [["mijoz", "summa"], ["Akmal", 5]])
    xlsm_path = _with_entries(
        _xlsx(tmp_path / "makros.xlsm", [["Vali", 7]]), {"xl/vbaProject.bin": b"v"}
    )
    for path, needle, kind in (
        (pdf, "Hisob 5000 USD", "ok"),
        (docx_path, "Yuk tashish shartnomasi", "ok"),
        (xlsx_path, "Akmal | 5", "ok"),
        (xlsm_path, "Vali | 7", "macro"),
    ):
        assert plan(filename=path.name).read_document is True
        verdict = safety.sniff_magic(path, path.suffix)
        assert (verdict.kind, verdict.readable) == (kind, True), path.name
        assert needle in documents.read_document(path).text, path.name


# --- the bot --------------------------------------------------------------------


@pytest.fixture
def media_in_tmp(monkeypatch, tmp_path):
    media = tmp_path / "media"
    media.mkdir()
    monkeypatch.setattr(type(settings), "media_dir", property(lambda self: media))
    return media


@pytest.fixture
def bot_stubs(monkeypatch):
    state = SimpleNamespace(downloads=[], content=b"", extracted=[])

    async def _download(bot, message, media, path):
        state.downloads.append(path)
        path.write_bytes(state.content)
        return True

    async def _process(session, interaction):
        state.extracted.append(interaction.id)
        interaction.processed = True
        return SimpleNamespace(ok=True, applied=Applied())

    async def _typing(message):
        return None

    monkeypatch.setattr(handlers, "_download", _download)
    monkeypatch.setattr(handlers, "_typing", _typing)
    monkeypatch.setattr(handlers, "process_interaction", _process)
    return state


class _DocMessage(_Message):
    def __init__(self, filename: str, *, mime="application/octet-stream", caption=None):
        super().__init__()
        self.document = SimpleNamespace(
            file_name=filename, file_size=2048, mime_type=mime
        )
        self.caption = caption
        self.date = datetime.now(TZ)
        self.message_id = 41
        self.reply_to_message = None
        for attr in ("video", "video_note", "voice", "audio", "photo"):
            setattr(self, attr, None)


async def _interactions(session) -> list[m.Interaction]:
    return list(await session.scalars(sa.select(m.Interaction)))


@pytest.mark.parametrize("name", ["ilova.apk", "Rasmlar.jpg.apk", "arxiv.zip"])
async def test_bot_document_handler_never_downloads_dangerous(
    bound,  # noqa: F811
    bot_stubs,
    media_in_tmp,
    name,
):
    message = _DocMessage(name)
    await handlers.on_document(message, bot=None)

    assert bot_stubs.downloads == []
    assert list(media_in_tmp.iterdir()) == []
    android = name.endswith(".apk")
    assert message.sent == [(replies.dangerous_file_refused(name, android=android), None)]
    assert await _interactions(bound) == []


async def test_the_refusal_says_what_the_plan_says(bound, bot_stubs, media_in_tmp):  # noqa: F811
    message = _DocMessage("ilova.apk")
    await handlers.on_document(message, bot=None)
    [(text, _)] = message.sent
    assert text == (
        "⚠️ «ilova.apk» — o'rnatiladigan dastur yoki arxiv. MIYA bunday fayllarni "
        "yuklab olmaydi va ochmaydi. Ichidagi ma'lumot kerak bo'lsa, PDF yoki rasm "
        "qilib yuboring.\n"
        "MIYA ilovasi faqat MIYA botidan keladi; boshqa hech kimdan kelgan APK faylni "
        "o'rnatmang."
    )


async def test_process_and_client_import_refuse_before_downloading(
    bound,  # noqa: F811
    bot_stubs,
    media_in_tmp,
):
    target = _DocMessage("yuk.rar")
    command = _DocMessage("x")
    command.document = None
    command.reply_to_message = target
    await handlers.cmd_process(command, bot=None)

    imported = _DocMessage("mijozlar.csv.exe", caption="/kodlar")
    await handlers.cmd_code_import(imported, bot=None)

    assert bot_stubs.downloads == []
    assert list(media_in_tmp.iterdir()) == []
    assert command.sent == [(replies.dangerous_file_refused("yuk.rar"), None)]
    assert imported.sent == [(replies.dangerous_file_refused("mijozlar.csv.exe"), None)]
    assert await _interactions(bound) == []


async def test_a_disguised_installer_is_deleted_before_parsing(
    bound,  # noqa: F811
    bot_stubs,
    media_in_tmp,
):
    bot_stubs.content = _windows_program()
    message = _DocMessage("hisob.pdf")
    await handlers.on_document(message, bot=None)

    [path] = bot_stubs.downloads
    assert not path.exists()
    assert message.sent == [(replies.dangerous_file_refused("hisob.pdf"), None)]
    assert await _interactions(bound) == []


async def test_bot_reads_allowed_documents_and_warns_about_macros(
    bound,  # noqa: F811
    bot_stubs,
    media_in_tmp,
    tmp_path,
):
    bot_stubs.content = _pdf("Hisob 5000 USD")
    pdf = _DocMessage("hisob.pdf")
    await handlers.on_document(pdf, bot=None)

    source = _with_entries(
        _docx(tmp_path / "x.docx", "Makrosli shartnoma"), {"word/vbaProject.bin": b"v"}
    )
    bot_stubs.content = source.read_bytes()
    macro = _DocMessage("shartnoma.docx")
    await handlers.on_document(macro, bot=None)

    first, second = sorted(await _interactions(bound), key=lambda i: i.id)
    assert "Hisob 5000 USD" in first.transcript
    assert "risk" not in first.media
    assert replies.macro_document_note("hisob.pdf") not in pdf.sent[0][0]

    assert "Makrosli shartnoma" in second.transcript
    assert second.media["risk"] == {"kind": "macro", "name": "shartnoma.docx"}
    [(text, _)] = macro.sent
    assert text.endswith(
        "ℹ️ «shartnoma.docx» makrosli hujjat. MIYA uni faqat matn sifatida o'qidi. "
        "Kompyuterda ochsangiz, makrosni yoqmang."
    )


async def test_a_mismatched_document_is_kept_but_never_parsed(
    bound,  # noqa: F811
    bot_stubs,
    media_in_tmp,
    monkeypatch,
):
    def _parse(path):
        raise AssertionError("a mismatched file must not reach a parser")

    monkeypatch.setattr(documents, "read_document", _parse)
    bot_stubs.content = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(64)
    message = _DocMessage("narxlar.xlsx")
    await handlers.on_document(message, bot=None)

    [row] = await _interactions(bound)
    assert row.needs_review is True
    assert message.sent == [(replies.DOCUMENT_FAILED_HINT, None)]


async def test_a_big_spreadsheet_sent_to_the_bot_is_kept_not_refused(
    bound,  # noqa: F811
    bot_stubs,
    media_in_tmp,
    monkeypatch,
    big_xlsx,
):
    def _parse(path):
        raise AssertionError("an over-limit container must not reach a parser")

    monkeypatch.setattr(documents, "read_document", _parse)
    bot_stubs.content = big_xlsx.read_bytes()
    message = _DocMessage("hisobot.xlsx")
    await handlers.on_document(message, bot=None)

    [path] = bot_stubs.downloads
    assert path.exists()
    [row] = await _interactions(bound)
    assert row.needs_review is True
    assert row.media["path"] == str(path)
    assert message.sent == [(replies.DOCUMENT_FAILED_HINT, None)]


async def test_an_unreadable_client_list_keeps_its_row(
    bound,  # noqa: F811
    bot_stubs,
    media_in_tmp,
):
    """The download must stay reachable by retention through its row."""
    bot_stubs.content = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(64)
    message = _DocMessage("mijozlar.xlsx", caption="/kodlar")
    await handlers.cmd_code_import(message, bot=None)

    [path] = bot_stubs.downloads
    [row] = await _interactions(bound)
    assert row.media["path"] == str(path)
    assert row.meta == {"kind": "client_import"}
    assert message.sent == [(replies.IMPORT_BAD_FILE, None)]


class _AudioMessage(_DocMessage):
    def __init__(self, filename: str) -> None:
        super().__init__(filename)
        self.audio = self.document
        self.audio.duration = 3
        self.document = None


async def test_an_installer_named_like_a_song_is_never_transcribed(
    bound,  # noqa: F811
    bot_stubs,
    media_in_tmp,
    monkeypatch,
):
    async def _transcribe(session, interaction, path):
        raise AssertionError("an installer must not reach the transcriber")

    monkeypatch.setattr(handlers, "transcribe_into", _transcribe)
    bot_stubs.content = _windows_program()
    sent = _AudioMessage("musiqa.mp3")
    await handlers.on_voice(sent, bot=None)

    target = _AudioMessage("qoshiq.mp3")
    command = _DocMessage("x")
    command.document = None
    command.reply_to_message = target
    await handlers.cmd_process(command, bot=None)

    assert len(bot_stubs.downloads) == 2
    assert not any(path.exists() for path in bot_stubs.downloads)
    assert sent.sent == [(replies.dangerous_file_refused("musiqa.mp3"), None)]
    assert command.sent == [(replies.dangerous_file_refused("qoshiq.mp3"), None)]
    assert await _interactions(bound) == []


async def test_the_refusal_shows_a_spoofed_name_as_it_really_is(
    bound,  # noqa: F811
    bot_stubs,
    media_in_tmp,
):
    message = _DocMessage("Rasm\u202egpj\u200b.apk")
    await handlers.on_document(message, bot=None)
    [(text, _)] = message.sent
    assert text.startswith("⚠️ «Rasmgpj.apk» — o'rnatiladigan dastur yoki arxiv.")
    assert "\u202e" not in text and "\u200b" not in text


# --- the userbot ----------------------------------------------------------------


class _Client:
    def __init__(self, content: bytes = b"") -> None:
        self.content = content
        self.downloads: list[str] = []

    async def download_media(self, message, file):
        self.downloads.append(file)
        Path(file).write_bytes(self.content)


def _tg_document(
    mid: int,
    name: str,
    *,
    size: int = 2048,
    out: bool = False,
    at: datetime | None = None,
    mime: str = "application/octet-stream",
):
    async def get_chat():
        return PEER

    async def get_sender():
        return PEER

    return SimpleNamespace(
        id=mid,
        chat_id=CHAT,
        date=at or datetime.now(TZ) - timedelta(minutes=1),
        message="",
        out=out,
        mentioned=False,
        file=SimpleNamespace(name=name, size=size, mime_type=mime),
        sticker=None,
        gif=None,
        voice=None,
        audio=None,
        video_note=None,
        photo=None,
        video=None,
        document=object(),
        get_chat=get_chat,
        get_sender=get_sender,
    )


async def _private_monitor() -> None:
    async with session_scope() as session:
        session.add(
            m.ChatMonitor(
                tg_chat_id=CHAT,
                chat_type=ChatType.private,
                title="Akmal",
                monitor_enabled=True,
                monitoring_since=datetime.now(TZ) - timedelta(days=30),
            )
        )


async def _row(session, mid: int) -> m.Interaction:
    rows = await session.scalars(
        sa.select(m.Interaction).where(m.Interaction.tg_chat_id == CHAT)
    )
    [row] = [r for r in rows if r.meta["tg_message_id"] == mid]
    await session.refresh(row)
    return row


async def test_ingest_never_fetches_a_large_apk_and_queues_one_warning(
    session, media_in_tmp
):
    await _private_monitor()
    client = _Client()
    message = _tg_document(7, "WhatsApp_Pro.apk", size=settings.doc_max_bytes * 3)
    assert await userbot.ingest_message(client, message) is True

    row = await _row(session, 7)
    assert client.downloads == []
    assert "approval" not in row.media
    assert row.media["skipped"] == "dangerous_file"
    assert row.media["risk"] == {"kind": "exec", "name": "WhatsApp_Pro.apk"}
    assert row.meta["safety"] == {
        "verdict": "dangerous_file",
        "score": 100,
        "reasons": ["exec_file", "suffix"],
    }
    assert row.meta["safety_notice"] == "dangerous_file"
    assert await approvals.awaiting_question(session) == []


async def test_the_owners_own_files_and_old_history_queue_no_warning(
    session, media_in_tmp
):
    await _private_monitor()
    client = _Client()
    await userbot.ingest_message(client, _tg_document(8, "arxiv.zip", out=True))
    old = datetime.now(TZ) - safety.NOTICE_MAX_AGE - timedelta(hours=1)
    await userbot.ingest_message(client, _tg_document(9, "ilova.apk", at=old))
    await userbot.ingest_message(client, _tg_document(10, "x.apk"), archive=True)

    for mid in (8, 9, 10):
        row = await _row(session, mid)
        assert row.media["skipped"] == "dangerous_file", mid
        assert row.meta["safety"]["verdict"] == "dangerous_file", mid
        assert "safety_notice" not in row.meta, mid
    assert client.downloads == []


async def test_ingest_deletes_an_installer_named_like_a_pdf(session, media_in_tmp):
    await _private_monitor()
    client = _Client(_windows_program())
    await userbot.ingest_message(client, _tg_document(11, "hisob.pdf"))

    row = await _row(session, 11)
    [path] = client.downloads
    assert not Path(path).exists()
    assert "path" not in row.media
    assert row.media["skipped"] == "dangerous_file"
    assert row.media["risk"] == {"kind": "exec", "name": "hisob.pdf"}
    assert row.meta["safety_notice"] == "dangerous_file"
    assert row.transcript is None


async def test_ingest_reads_a_macro_document_as_text_and_says_so(
    session, media_in_tmp, tmp_path
):
    await _private_monitor()
    source = _with_entries(
        _xlsx(tmp_path / "s.xlsm", [["Akmal", 5]]), {"xl/vbaProject.bin": b"v"}
    )
    client = _Client(source.read_bytes())
    await userbot.ingest_message(client, _tg_document(12, "hisob.xlsm"))

    row = await _row(session, 12)
    assert "Akmal | 5" in row.transcript
    assert row.media["risk"]["kind"] == "macro"
    assert row.meta["safety_notice"] == "macro"
    assert "safety" not in row.meta


async def test_ingest_never_parses_a_file_that_is_not_what_it_claims(
    session, media_in_tmp, monkeypatch
):
    def _parse(path):
        raise AssertionError("a mismatched file must not reach a parser")

    monkeypatch.setattr(documents, "read_document", _parse)
    await _private_monitor()
    client = _Client(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(64))
    await userbot.ingest_message(client, _tg_document(15, "narxlar.xlsx"))

    row = await _row(session, 15)
    assert row.media["skipped"] == "mismatched_type"
    assert row.needs_review is True
    assert row.transcript is None
    assert "safety_notice" not in row.meta


async def test_ingest_still_reads_an_ordinary_pdf(session, media_in_tmp):
    await _private_monitor()
    client = _Client(_pdf("Konteyner GS367"))
    await userbot.ingest_message(client, _tg_document(13, "yuk.pdf"))

    row = await _row(session, 13)
    assert "Konteyner GS367" in row.transcript
    assert "risk" not in row.media and "safety_notice" not in row.meta


async def test_ingest_never_transcribes_an_installer_named_like_a_song(
    session, media_in_tmp, monkeypatch
):
    class _Scribe:
        async def transcribe(self, path):
            raise AssertionError("an installer must not reach the transcriber")

    monkeypatch.setattr(userbot, "get_transcriber", lambda: _Scribe())
    await _private_monitor()
    client = _Client(_windows_program())
    message = _tg_document(16, "musiqa.mp3", mime="audio/mpeg")
    message.audio = object()
    await userbot.ingest_message(client, message)

    row = await _row(session, 16)
    [path] = client.downloads
    assert not Path(path).exists()
    assert "path" not in row.media
    assert row.media["skipped"] == "dangerous_file"
    assert row.media["risk"] == {"kind": "exec", "name": "musiqa.mp3"}
    assert row.meta["safety_notice"] == "dangerous_file"
    assert row.transcript is None


async def test_ingest_keeps_a_big_spreadsheet_without_a_warning(
    session, media_in_tmp, monkeypatch, big_xlsx
):
    def _parse(path):
        raise AssertionError("an over-limit container must not reach a parser")

    monkeypatch.setattr(documents, "read_document", _parse)
    await _private_monitor()
    client = _Client(big_xlsx.read_bytes())
    size = big_xlsx.stat().st_size
    await userbot.ingest_message(client, _tg_document(17, "hisobot.xlsx", size=size))

    row = await _row(session, 17)
    [path] = client.downloads
    assert Path(path).exists() and row.media["path"] == path
    assert row.media["skipped"] == "mismatched_type"
    assert row.needs_review is True
    assert "risk" not in row.media
    assert "safety" not in row.meta and "safety_notice" not in row.meta


# --- the worker -----------------------------------------------------------------


class _Bot:
    def __init__(self, *, fail: bool = False) -> None:
        self.sent: list[tuple[str, dict]] = []
        self.fail = fail

    async def send_message(self, chat_id, text, **kwargs):
        if self.fail:
            raise RuntimeError("telegram down")
        self.sent.append((text, kwargs))


async def _queued(
    session,
    *,
    name="ilova.apk",
    chat_type=ChatType.private,
    title="Akmal",
    notice="dangerous_file",
    at=DAY,
) -> m.Interaction:
    person = m.Person(display_name="Akmal", aliases=[])
    session.add(person)
    session.add(m.ChatMonitor(tg_chat_id=CHAT, chat_type=chat_type, title=title))
    await session.flush()
    row = m.Interaction(
        source=InteractionSource.telegram_userbot,
        direction=Direction.in_,
        person_id=person.id,
        tg_chat_id=CHAT,
        occurred_at=at,
        media={
            "type": "document",
            "filename": name,
            "risk": {"kind": "exec", "name": name},
        },
        meta={"tg_message_id": 1, "safety_notice": notice},
    )
    session.add(row)
    await session.commit()
    return row


async def test_warning_sent_once(session):
    row = await _queued(session)
    bot = _Bot()
    await worker.safety_notice_job(bot, now=DAY + timedelta(minutes=1))
    await worker.safety_notice_job(bot, now=DAY + timedelta(minutes=2))

    [(text, kwargs)] = bot.sent
    assert text == (
        "⚠️ Akmal sizga «ilova.apk» faylini yubordi (Telegram). Bu o'rnatiladigan "
        "dastur yoki arxiv — xavfli bo'lishi mumkin. Ochmang va o'rnatmang. MIYA uni "
        "yuklab olmadi.\n"
        "Tanish odamdan kelgan bo'lsa ham ehtiyot bo'ling. Shubha bo'lsa, unga "
        "qo'ng'iroq qilib so'rang.\n"
        "MIYA ilovasi faqat MIYA botidan keladi; boshqa hech kimdan kelgan APK faylni "
        "o'rnatmang."
    )
    assert kwargs["disable_notification"] is False
    await session.refresh(row)
    assert "safety_notified" in row.meta
    # A warning is not a tap-request: the day's question budget is untouched.
    assert await questions.spent_today(session, now=DAY + timedelta(minutes=2)) == 0
    assert await session.scalar(sa.select(sa.func.count(m.QuestionLog.id))) == 0


async def test_a_group_warning_names_the_group(session):
    await _queued(session, name="yuk.rar", chat_type=ChatType.group, title="Yuk guruhi")
    bot = _Bot()
    await worker.safety_notice_job(bot, now=DAY + timedelta(minutes=1))
    [(text, _)] = bot.sent
    assert text.startswith("⚠️ Akmal sizga «yuk.rar» faylini yubordi (Yuk guruhi).")
    assert replies.APK_RULE_LINE not in text


async def test_a_warning_is_silent_at_night_not_held(session):
    night = DAY.replace(hour=23, minute=45)
    await _queued(session, at=night)
    bot = _Bot()
    await worker.safety_notice_job(bot, now=night + timedelta(minutes=1))
    [(_, kwargs)] = bot.sent
    assert kwargs["disable_notification"] is True


async def test_an_undelivered_warning_stays_queued(session):
    row = await _queued(session)
    await worker.safety_notice_job(_Bot(fail=True), now=DAY + timedelta(minutes=1))
    await session.refresh(row)
    assert "safety_notified" not in row.meta

    bot = _Bot()
    await worker.safety_notice_job(bot, now=DAY + timedelta(minutes=2))
    assert len(bot.sent) == 1


async def test_a_macro_notice_uses_the_macro_line(session):
    await _queued(session, name="hisob.xlsm", notice="macro")
    bot = _Bot()
    await worker.safety_notice_job(bot, now=DAY + timedelta(minutes=1))
    assert [text for text, _ in bot.sent] == [replies.macro_document_note("hisob.xlsm")]


def test_the_job_is_registered_every_minute():
    """1 min, one at a time, missed runs coalesced (JV-01 behaviour 8)."""
    source = Path(worker.__file__).read_text()
    start = source.index("        safety_notice_job,\n")
    block = source[start : source.index("scheduler.add_job", start)]
    for needle in (
        "IntervalTrigger(minutes=1)",
        'id="safety_notices"',
        "max_instances=1",
        "coalesce=True",
    ):
        assert needle in block, needle


async def test_an_approved_installer_is_still_never_fetched(session, media_in_tmp):
    """A yes given before the shield existed cannot bring an installer over."""
    session.add(
        m.Interaction(
            source=InteractionSource.telegram_userbot,
            direction=Direction.in_,
            tg_chat_id=CHAT,
            occurred_at=datetime.now(TZ) - timedelta(minutes=5),
            media={
                "type": "document",
                "filename": "ilova.apk",
                "size": 90_000_000,
                "processed": False,
                "approval": {"state": approvals.APPROVED, "reason": "too_large"},
            },
            meta={"tg_message_id": 14},
        )
    )
    await session.commit()

    class _ApprovedClient(_Client):
        async def get_messages(self, chat_id, ids):
            return _tg_document(ids, "ilova.apk")

    client = _ApprovedClient()
    assert await userbot.fetch_approved(client) == 1

    row = await _row(session, 14)
    assert client.downloads == []
    assert row.media["skipped"] == "dangerous_file"
    assert row.media["risk"] == {"kind": "exec", "name": "ilova.apk"}
    assert approvals.state_of(row) == approvals.DONE
    assert row.meta["safety_notice"] == "dangerous_file"


async def test_a_warning_shows_a_spoofed_name_as_it_really_is(session):
    await _queued(session, name="Rasm\u202egpj\u200b.apk")
    bot = _Bot()
    await worker.safety_notice_job(bot, now=DAY + timedelta(minutes=1))
    [(text, _)] = bot.sent
    assert text.startswith("⚠️ Akmal sizga «Rasmgpj.apk» faylini yubordi (Telegram).")
    assert "\u202e" not in text and "\u200b" not in text


async def test_an_installer_queued_before_the_shield_is_never_offered(session):
    """The size check used to run first, so a big .apk could be waiting for a
    download question when the shield was deployed."""
    stale = m.Interaction(
        source=InteractionSource.telegram_userbot,
        direction=Direction.in_,
        tg_chat_id=CHAT,
        occurred_at=datetime.now(TZ) - timedelta(hours=2),
        media={
            "type": "document",
            "filename": "ilova.apk",
            "size": 60_000_000,
            "processed": False,
            "approval": {"state": approvals.PENDING, "reason": "too_large"},
        },
        meta={"tg_message_id": 18},
    )
    session.add(stale)
    await session.commit()

    for for_push in (True, False):
        items = await questions.collect(session, now=datetime.now(TZ), for_push=for_push)
        assert [i for i in items if i.kind == questions.KIND_MEDIA] == [], for_push
    await session.commit()

    await session.refresh(stale)
    assert approvals.state_of(stale) == approvals.DONE
    assert stale.media["approval"]["error"] == "dangerous_file"
    assert stale.media["skipped"] == "dangerous_file"
    assert stale.media["risk"] == {"kind": "exec", "name": "ilova.apk"}
    assert stale.meta["safety"]["verdict"] == "dangerous_file"
    assert stale.meta["safety_notice"] == "dangerous_file"
    assert await approvals.awaiting_question(session) == []
