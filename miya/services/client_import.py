"""Import the owner's existing client-code list (WP-34).

A forwarding business already keeps "GS367 — Akmal" somewhere, usually in
Excel. ``read_rows`` reads a CSV or .xlsx, ``plan_import`` sorts every row
into what would happen without writing anything, and ``apply_import``
writes the attach and create rows in the caller's transaction. A code held
by someone else is never changed: that is a conflict for the owner to see.
Running the same file twice changes nothing the second time.
"""

from __future__ import annotations

import csv
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import sqlalchemy as sa
from sqlalchemy.ext.asyncio import AsyncSession

from miya.db.models import Person
from miya.services import codes
from miya.services.people import (
    MATCH_THRESHOLD,
    QUESTION_THRESHOLD,
    best_match,
    find_by_phone,
    find_person,
    normalise,
)
from miya.services.text import fold_apostrophes

MAX_ROWS = 5000
DEFAULT_RELATIONSHIP = "mijoz"

_HEADERS = {
    "code": {"kod", "code", "gs", "mijoz kodi"},
    "name": {"ism", "name", "mijoz", "fio", "ф.и.о"},
    "phone": {"telefon", "phone", "tel", "raqam"},
    "telegram": {"telegram", "username", "tg"},
    "note": {"izoh", "note", "relationship", "kim"},
}


class BadFile(Exception):
    """Not a CSV or .xlsx, or nothing in it that reads as a client list."""


@dataclass(slots=True)
class Row:
    code_raw: str
    name: str
    phone: str
    telegram: str
    note: str
    line_no: int


@dataclass(slots=True)
class Planned:
    row: Row
    code: str | None = None
    person: Person | None = None
    # Who the code is shown against in a conflict: the holder in the
    # database, or the other name the same code carries in the file.
    against: str = ""


@dataclass(slots=True)
class ImportPlan:
    attach: list[Planned] = field(default_factory=list)
    create: list[Planned] = field(default_factory=list)
    same: list[Planned] = field(default_factory=list)
    conflict: list[Planned] = field(default_factory=list)
    bad: list[Planned] = field(default_factory=list)

    def counts(self) -> dict[str, int]:
        return {
            "attach": len(self.attach),
            "create": len(self.create),
            "same": len(self.same),
            "conflict": len(self.conflict),
            "bad": len(self.bad),
        }


# --- reading -------------------------------------------------------------------


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _header_key(value: str) -> str | None:
    folded = fold_apostrophes(value).strip().lower().replace(".", "")
    for key, names in _HEADERS.items():
        if folded in {n.replace(".", "") for n in names}:
            return key
    return None


def _table(path: Path) -> list[list[str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv" or suffix == ".txt":
        try:
            text = path.read_text(encoding="utf-8-sig")
        except (OSError, UnicodeDecodeError) as exc:
            raise BadFile(str(exc)) from exc
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [[_cell(c) for c in row] for row in csv.reader(text.splitlines(), dialect)]
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        try:
            book = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many kinds for a bad file
            raise BadFile(str(exc)) from exc
        try:
            sheet = book.worksheets[0]
            return [[_cell(c) for c in row] for row in sheet.iter_rows(values_only=True)]
        finally:
            book.close()
    raise BadFile(f"unsupported file type: {suffix or '(none)'}")


def read_rows(path: str | Path) -> list[Row]:
    """The rows of a client list; raises BadFile when there are none."""
    table = _table(Path(path))
    if not table:
        raise BadFile("empty file")
    columns: dict[str, int] = {}
    first = next((i for i, row in enumerate(table) if any(row)), None)
    if first is None:
        raise BadFile("empty file")
    for index, value in enumerate(table[first]):
        key = _header_key(value)
        if key is not None and key not in columns:
            columns[key] = index
    if columns:
        if "code" not in columns:
            raise BadFile("no code column")
        body = table[first + 1 :]
        start = first + 2
    else:
        columns = {"code": 0, "name": 1}
        body = table[first:]
        start = first + 1

    def get(row: list[str], key: str) -> str:
        index = columns.get(key)
        return row[index] if index is not None and index < len(row) else ""

    rows: list[Row] = []
    for offset, raw in enumerate(body):
        if not any(raw):
            continue
        rows.append(
            Row(
                code_raw=get(raw, "code"),
                name=get(raw, "name"),
                phone=get(raw, "phone"),
                telegram=get(raw, "telegram"),
                note=get(raw, "note"),
                line_no=start + offset,
            )
        )
        if len(rows) >= MAX_ROWS:
            break
    if not rows:
        raise BadFile("no rows")
    return rows


# --- planning ------------------------------------------------------------------


def _digits(phone: str) -> str:
    return re.sub(r"\D", "", phone or "")


async def plan_import(session: AsyncSession, rows: list[Row]) -> ImportPlan:
    """What an import would do, row by row. Writes nothing."""
    plan = ImportPlan()
    names_in_file = Counter(normalise(r.name) for r in rows if r.name.strip())
    seen: dict[str, Planned] = {}
    clashing: set[str] = set()
    candidates: list[Planned] = []

    for row in rows:
        code = codes.canonical_client_code(row.code_raw)
        if code is None or not (row.name.strip() or len(_digits(row.phone)) >= 7):
            plan.bad.append(Planned(row=row, code=code))
            continue
        planned = Planned(row=row, code=code)
        earlier = seen.get(code)
        if earlier is not None:
            if normalise(earlier.row.name) == normalise(row.name):
                continue  # the same line twice
            earlier.against = row.name
            planned.against = earlier.row.name
            clashing.add(code)
        else:
            seen[code] = planned
        candidates.append(planned)

    for planned in candidates:
        row, code = planned.row, planned.code
        if code in clashing:
            plan.conflict.append(planned)
            continue
        holder = await codes.holder(session, code)
        if holder is not None:
            planned.person = holder
            if not row.name.strip() or best_match(row.name, [holder])[1] >= (
                QUESTION_THRESHOLD
            ):
                plan.same.append(planned)
            else:
                planned.against = holder.display_name
                plan.conflict.append(planned)
            continue
        person = await _existing(session, row, names_in_file)
        if person is not None:
            planned.person = person
            plan.attach.append(planned)
        else:
            plan.create.append(planned)
    return plan


async def _existing(
    session: AsyncSession, row: Row, names_in_file: Counter
) -> Person | None:
    if len(_digits(row.phone)) >= 7:
        by_phone = await find_by_phone(session, row.phone)
        if by_phone is not None:
            return by_phone
    username = row.telegram.strip().lstrip("@")
    if username:
        by_username = await session.scalar(
            sa.select(Person).where(
                sa.func.lower(Person.telegram_username) == username.lower()
            )
        )
        if by_username is not None:
            return by_username
    if not row.name.strip() or names_in_file[normalise(row.name)] > 1:
        return None
    match = await find_person(session, row.name)
    if match.person is None or match.ambiguous:
        return None
    if not (match.exact or match.score >= MATCH_THRESHOLD):
        return None
    if await codes.codes_of(session, match.person.id):
        return None  # a namesake who is already another client
    return match.person


# --- applying ------------------------------------------------------------------


def _fill(person: Person, row: Row) -> None:
    digits = _digits(row.phone)
    if len(digits) >= 7 and not person.phone:
        person.phone = digits
    username = row.telegram.strip().lstrip("@")
    if username and not person.telegram_username:
        person.telegram_username = username


async def apply_import(session: AsyncSession, plan: ImportPlan, *, by: str) -> dict:
    """Write the attach and create rows; conflicts are never touched."""
    written = {"attach": 0, "create": 0}
    for planned in plan.attach:
        _fill(planned.person, planned.row)
        try:
            await codes.attach(
                session, planned.person, planned.code, source="import", by=by
            )
        except codes.CodeTaken:
            continue
        written["attach"] += 1
    for planned in plan.create:
        row = planned.row
        digits = _digits(row.phone)
        person = Person(
            display_name=row.name.strip() or planned.code,
            aliases=[],
            phone=digits if len(digits) >= 7 else None,
            telegram_username=row.telegram.strip().lstrip("@") or None,
            relationship_=row.note.strip() or DEFAULT_RELATIONSHIP,
        )
        session.add(person)
        await session.flush()
        try:
            await codes.attach(session, person, planned.code, source="import", by=by)
        except codes.CodeTaken:
            await session.delete(person)
            continue
        planned.person = person
        written["create"] += 1
    await session.flush()
    return written
